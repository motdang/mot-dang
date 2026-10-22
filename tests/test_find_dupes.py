#!/usr/bin/env python3
"""find_dupes: which made-up records pair up, and what a labeled sheet writes back.

  1. phone numbers: formats that are one number, short codes left out
  2. pairs: phone within 150 m, name within 60 m, a wider reach for a rare name,
     numbered names that differ stay apart, other provinces stay apart
  3. pairs already joined or called different are not listed again
  4. the sheet: header, dropdown, frozen panes; the answers read back
  5. --apply: same -> merges.json (three-way and older merges flattened),
     different -> not_same.json, the ledger marks the pair

    python3 tests/test_find_dupes.py
"""
import json
import sys
import tempfile
from pathlib import Path
from types import SimpleNamespace

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "importers"))
import find_dupes as F  # noqa: E402

failures = []


def check(label, ok, detail=""):
    print(f"  {'ok  ' if ok else 'FAIL'}  {label}{'' if ok else ' — ' + str(detail)}")
    if not ok:
        failures.append(label)


def rec(i, name, lat=18.79, lng=98.98, phone=None, en=None, prov="cm", cat=("food",), **kw):
    r = {"id": i, "province": prov, "name": name, "nameEn": en, "lat": lat, "lng": lng,
         "phone": phone, "cat": list(cat), "sub": [], "geoPrecision": "exact",
         "attrs": {}, "sources": [{"type": "osm"}]}
    r.update(kw)
    return r


# 1. phone numbers
check("+66 and 0 forms are one number",
      F.phone_keys("+66871824663") == F.phone_keys("087-182-4663") == {"871824663"})
check("two numbers in one string", len(F.phone_keys("053-221-770,053-221-771")) == 2)
check("landline with and without the zero", F.phone_keys("+6653221770") == F.phone_keys("053221770"))
check("short code left out", F.phone_keys("1150") == set())

# 2. pairs
M = 0.00001  # about 1.1 m of latitude
recs = [F.slim(r) for r in [
    rec("cm-a1", "Blue Door Cafe", phone="+66811111111"),
    rec("cm-a2", "ร้านอื่น", lat=18.79 + 100 * M, phone="081-111-1111"),        # phone, 111 m
    rec("cm-a3", "Far Shop", lat=18.79 + 200 * M, phone="081-111-1111"),        # phone, 222 m: too far
    rec("cm-b1", "Green Hill Resort", lng=98.99),
    rec("cm-b2", "green hill resort!", lng=98.99, lat=18.79 + 30 * M),          # name, 33 m
    rec("cm-c1", "Tower 7", lng=99.0),
    rec("cm-c2", "Tower 9", lng=99.0),                                          # numbers differ
    rec("cm-d1", "Gate Moat Old", lng=99.01),
    rec("cm-d2", "Gate Moat Old", lng=99.01, lat=18.79 + 90 * M),               # rare name, 100 m
    rec("cr-e1", "Same Name Place", lng=99.02, prov="cr"),
    rec("cm-e2", "Same Name Place", lng=99.02),                                 # other province
    rec("cm-f1", "Maetaeng Elephant Park", lng=99.03),
    rec("cm-f2", "Maetaeng Elephant Park Camp", lng=99.03, lat=18.79 + 20 * M),  # similar
    rec("cm-g1", "Pinless", lat=None, lng=None, phone="082-222-2222"),
    rec("cm-g2", "Pinless", lat=None, lng=None, phone="082-222-2222"),
]]
pairs, skipped, pinned, put = F.find_pairs(recs, set())
like = F.Likeness(ROOT)
F.find_similar(pairs, pinned, like, put)
P = lambda a, b: pairs.get(F.pair_id(a, b))  # noqa: E731
check("phone pair within 150 m", P("cm-a1", "cm-a2") and "phone" in P("cm-a1", "cm-a2")["why"])
check("phone pair beyond 150 m is not listed", P("cm-a1", "cm-a3") is None)
check("same name within 60 m", P("cm-b1", "cm-b2") and "name" in P("cm-b1", "cm-b2")["why"])
check("rare name reaches 120 m", P("cm-d1", "cm-d2") is not None)
check("numbered names that differ stay apart", P("cm-c1", "cm-c2") is None)
check("pairs never cross provinces", P("cr-e1", "cm-e2") is None)
check("similar name within 40 m", like.ok and P("cm-f1", "cm-f2") and "similar" in P("cm-f1", "cm-f2")["why"])
check("records without a pin never pair", P("cm-g1", "cm-g2") is None)

# ranking and keeper
paths = {}
e1, e2 = P("cm-b1", "cm-b2"), P("cm-a1", "cm-a2")
check("a name beats nothing", F.score(e1, paths)[0] > 0)
check("keeper is the fuller record",
      F.pick_keeper(F.slim(rec("x1", "A", phone="1234567890", en="A")),
                    F.slim(rec("x2", "A")), {}, {}) == "A")

# 3. known pairs
skip = {F.pair_id("cm-b1", "cm-b2")}
pairs2, _, _, _ = F.find_pairs(recs, skip)
check("a skipped pair is not listed", F.pair_id("cm-b1", "cm-b2") not in pairs2)

# 4 + 5. sheet and apply, on a scratch data root
with tempfile.TemporaryDirectory() as tmp:
    root = Path(tmp)
    (root / "data" / "canonical").mkdir(parents=True)
    (root / "data" / "curated").mkdir(parents=True)
    ids = ["cm-k1", "cm-k2", "cm-k3", "cm-m1", "cm-m2", "cm-n1", "cm-n2", "cm-old", "cm-oldchild"]
    (root / "data" / "canonical" / "cm.json").write_text(json.dumps([rec(i, i) for i in ids]))
    (root / "data" / "curated" / "merges.json").write_text(json.dumps(
        {"_readme": "r", "merges": {"cm-old": ["cm-oldchild"]}, "_why": {}}))
    rows = []
    for a, b in (("cm-k1", "cm-k2"), ("cm-k2", "cm-k3"), ("cm-m1", "cm-m2"), ("cm-n1", "cm-n2"),
                 ("cm-k1", "cm-old")):
        ra, rb = F.slim(rec(a, a)), F.slim(rec(b, b))
        rows.append({"pid": F.pair_id(a, b), "a": ra, "b": rb, "d": 5.0, "pts": 90,
                     "band": F.band(90), "words": "w", "why": ["name"], "keep": "A",
                     "url_a": "https://motdang.net/cm/p/x.html", "url_b": "", "map": "m"})
    out = root / "_incoming" / "dupes"
    out.mkdir(parents=True)
    sheet = out / "dupes-test.xlsx"
    F.write_xlsx(sheet, rows)
    from openpyxl import load_workbook
    wb = load_workbook(sheet)
    ws = wb["Pairs"]
    check("header row frozen with the label column", ws.freeze_panes == "B2", ws.freeze_panes)
    check("dropdown on the label column", bool(ws.data_validations.dataValidation)
          and str(ws.data_validations.dataValidation[0].sqref) == "A2:A6")
    check("a page link is a link", ws["I2"].hyperlink is not None and ws["J2"].hyperlink is None)
    answers = [F.SAME, F.SAME, "ไม่แน่ใจ · Not sure", F.DIFF, "same"]
    for n, ans in enumerate(answers, 2):
        ws.cell(row=n, column=1).value = ans
    wb.save(sheet)
    got, unread = F.read_labels(sheet)
    check("answers read back", [g[0] for g in got] == ["same", "same", "unsure", "different", "same"], got)
    F.run_apply(SimpleNamespace(data=str(root), out=str(out), apply=str(sheet), by="Test", dry=False))
    m = json.loads((root / "data" / "curated" / "merges.json").read_text())["merges"]
    check("three-way group lands under one keeper",
          sorted(m.get("cm-k1", [])) == ["cm-k2", "cm-k3", "cm-old", "cm-oldchild"] or
          any(sorted(v + [k]) == sorted(["cm-k1", "cm-k2", "cm-k3", "cm-old", "cm-oldchild"])
              for k, v in m.items()), m)
    check("one merge entry per keeper, no drop listed twice",
          len([d for v in m.values() for d in v]) == len({d for v in m.values() for d in v}), m)
    check("not sure and different are not merged", "cm-m1" not in str(m) and "cm-n1" not in str(m), m)
    ns = json.loads((root / "data" / "curated" / "not_same.json").read_text())
    check("different pair recorded", ns["pairs"] == [["cm-n1", "cm-n2"]], ns["pairs"])
    ledger = F.read_ledger(out / "ledger.tsv")
    check("ledger marks the labels", ledger[F.pair_id("cm-n1", "cm-n2")]["label"] == "different"
          and ledger[F.pair_id("cm-m1", "cm-m2")]["label"] == "unsure")
    before = (root / "data" / "curated" / "merges.json").read_text()
    F.run_apply(SimpleNamespace(data=str(root), out=str(out), apply=str(sheet), by="Test", dry=False))
    check("a second apply changes nothing", (root / "data" / "curated" / "merges.json").read_text() == before)
    check("known pairs are skipped by the next run",
          F.pair_id("cm-n1", "cm-n2") in F.known_skips(root)
          and F.pair_id("cm-k2", "cm-k3") in F.known_skips(root))

print(f"\n{'FAILED: ' + ', '.join(failures) if failures else 'all passed'}")
sys.exit(1 if failures else 0)
